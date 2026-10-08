#!/usr/bin/env python3
"""Prüft die Unterrichtsmaterialien (materialien/**/*.docx, *.pptx, *.txt) auf Konventionsverstöße.

Fehler (Exit-Code 1):
  - Altbegriffe „Kata“, „Gilde(n)“, „XP“ (gültig: Warm-up, Team, Punkte im Lab)
  - generische Abiturklassen (Stack<…>, Queue<…>, DynArray<…>, BinTree<…>) im Code der Online-IDE-Dateien
  - nicht erlaubte Zeichenketten-Operationen (toLowerCase, split, indexOf, trim, replace …) im Code der Online-IDE-Dateien
Hinweise (kein Fehler):
  - „Array“ im Fließtext (Konvention: „Reihung“), sofern nicht als „statt/nie Array“ erklärt
  - System.out im Code der Online-IDE-Dateien (dort println ohne System.out)

Absichtliche Ausnahmen (z. B. Bug-Hunts, die genau diese Fehler zeigen) stehen in AUSNAHMEN.
Mit --rendern werden alle docx/pptx zusätzlich mit LibreOffice (soffice) nach PDF umgewandelt; gemeldet werden
Dateien, die sich nicht umwandeln lassen. Mit --pdf-ordner <dir> bleiben die PDFs zur Sichtprüfung erhalten.

Aufruf: python3 werkzeuge/materialien_pruefen.py [--rendern] [--pdf-ordner DIR] [--hinweise]
Benötigt python-docx und python-pptx (pip install python-docx python-pptx).
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

REPO = Path(__file__).resolve().parent.parent
MAT = REPO / "materialien"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

ALTBEGRIFFE = re.compile(r"\b(Kata|Katas|Gilde|Gilden|XP)\b")
GENERISCH = re.compile(r"\b(Stack|Queue|DynArray|BinTree)\s*<")
VERBOTEN = re.compile(r"\.(toLowerCase|toUpperCase|split|indexOf|lastIndexOf|trim|replace|replaceAll|startsWith|endsWith|"
                      r"equalsIgnoreCase|isBlank|strip|chars|toCharArray)\s*\(")
ARRAY = re.compile(r"\bArrays?\b")
ARRAY_OK = re.compile(r"Array-Literal|ArrayIndex|ArrayList|DynArray|statt\s+[„\"]?Array|nie\s+[„\"]?Array|"
                      r"Array\s+heißt|java\.util", re.I)

# (Datei-Muster, Prüfung, Begründung) – absichtlich so im Material
AUSNAHMEN = [
    (r"materialien/LIESMICH\.txt$", "altbegriff", "beschreibt das Upgrade von der alten WebApp (Gilden/XP)"),
    (r"Lehrerhinweise|Erwartungshorizont|_LSG", "altbegriff", "erklärt Begriffe der Vorgängerversion"),
]


def ausnahme(pfad, art):
    rel = str(pfad.relative_to(REPO))
    return any(re.search(m, rel) and a == art for m, a, _ in AUSNAHMEN)


def zeilen_docx(p):
    d = Document(p)
    teile = [d.element.body] + [s.header._element for s in d.sections] + [s.footer._element for s in d.sections]
    for teil in teile:
        for absatz in teil.iter(qn("w:p")):
            t = "".join(x.text or "" for x in absatz.iter(qn("w:t")))
            if t.strip():
                yield t


def zeilen_pptx(p):
    for i, folie in enumerate(Presentation(p).slides, 1):
        for form in folie.shapes:
            for absatz in form._element.iter(A_NS + "p"):
                t = "".join(x.text or "" for x in absatz.iter(A_NS + "t"))
                if t.strip():
                    yield f"Folie {i}: {t}"
        if folie.has_notes_slide:
            for t in folie.notes_slide.notes_text_frame.text.split("\n"):
                if t.strip():
                    yield f"Folie {i} (Notiz): {t}"


def code_zeilen(text):
    """Zeilen einer Online-IDE-Datei ohne Kommentare und ohne Zeichenkettenliterale (Meldungstexte)."""
    text = re.sub(r"/\*.*?\*/", lambda m: "\n" * m.group(0).count("\n"), text, flags=re.S)
    for nr, z in enumerate(text.split("\n"), 1):
        z = re.sub(r'"(?:\\.|[^"\\])*"', '""', z)
        z = z.split("//", 1)[0]
        yield nr, z


def pruefe(pfad):
    fehler, hinweise = [], []
    if pfad.suffix == ".txt":
        text = pfad.read_text(encoding="utf-8")
        zeilen = list(enumerate(text.split("\n"), 1))
        for nr, z in code_zeilen(text):
            if GENERISCH.search(z):
                fehler.append(f"Zeile {nr}: generische Abiturklasse ({GENERISCH.search(z).group(0)}…) – NI-Klassen sind nicht generisch")
            if VERBOTEN.search(z):
                fehler.append(f"Zeile {nr}: nicht erlaubte Zeichenketten-Operation {VERBOTEN.search(z).group(0)}")
            if "System.out" in z:
                hinweise.append(f"Zeile {nr}: System.out – in der Online-IDE println(...) ohne System.out")
        texte = [(f"Zeile {nr}", z) for nr, z in zeilen]
    elif pfad.suffix == ".docx":
        texte = [(f"Absatz {i}", t) for i, t in enumerate(zeilen_docx(pfad), 1)]
    else:
        texte = [("", t) for t in zeilen_pptx(pfad)]
    for ort, t in texte:
        m = ALTBEGRIFFE.search(t)
        if m and not ausnahme(pfad, "altbegriff"):
            fehler.append(f"{ort}: Altbegriff „{m.group(0)}“: {t.strip()[:120]}")
        if pfad.suffix != ".txt" and ARRAY.search(t) and not ARRAY_OK.search(t):
            hinweise.append(f"{ort}: „Array“ statt „Reihung“?: {t.strip()[:120]}")
    return fehler, hinweise


def rendern(dateien, pdf_ordner):
    """Wandelt alle Dateien mit soffice nach PDF; gibt {datei: problem} zurück."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("soffice (LibreOffice) nicht gefunden – --rendern nicht möglich.")
    probleme = {}
    with tempfile.TemporaryDirectory(prefix="lo-profil-") as profil:
        for p in dateien:
            ziel = Path(pdf_ordner) / p.parent.relative_to(MAT)
            ziel.mkdir(parents=True, exist_ok=True)
            r = subprocess.run([soffice, f"-env:UserInstallation=file://{profil}", "--headless", "--convert-to", "pdf",
                                "--outdir", str(ziel), str(p)], capture_output=True, text=True, timeout=300)
            pdf = ziel / (p.stem + ".pdf")
            if r.returncode or not pdf.exists() or pdf.stat().st_size < 1000:
                probleme[p] = (r.stderr or r.stdout).strip()[:200] or "keine PDF erzeugt"
    return probleme


def main():
    args = sys.argv[1:]
    mit_hinweisen = "--hinweise" in args
    dateien = sorted(p for p in MAT.rglob("*") if p.suffix in (".docx", ".pptx", ".txt") and not p.name.startswith("~$"))
    anzahl_fehler = anzahl_hinweise = 0
    for p in dateien:
        fehler, hinweise = pruefe(p)
        anzahl_fehler += len(fehler)
        anzahl_hinweise += len(hinweise)
        if fehler or (mit_hinweisen and hinweise):
            print(f"{'✘' if fehler else '·'} {p.relative_to(REPO)}")
            for f in fehler:
                print("    ! " + f)
            if mit_hinweisen:
                for h in hinweise:
                    print("    · " + h)
    if "--rendern" in args:
        pdf_ordner = args[args.index("--pdf-ordner") + 1] if "--pdf-ordner" in args else None
        office = [p for p in dateien if p.suffix in (".docx", ".pptx")]
        with tempfile.TemporaryDirectory(prefix="pdf-") as tmp:
            probleme = rendern(office, pdf_ordner or tmp)
        for p, grund in probleme.items():
            print(f"✘ {p.relative_to(REPO)}: lässt sich nicht nach PDF umwandeln ({grund})")
        anzahl_fehler += len(probleme)
        print(f"{len(office)} Office-Dateien nach PDF umgewandelt" + (f" (in {pdf_ordner})" if pdf_ordner else "") +
              f", {len(probleme)} Problem(e).")
    print(f"{len(dateien)} Dateien geprüft: {anzahl_fehler} Fehler, {anzahl_hinweise} Hinweis(e)"
          + ("" if mit_hinweisen else " (anzeigen mit --hinweise)") + ".")
    sys.exit(1 if anzahl_fehler else 0)


if __name__ == "__main__":
    main()
